
import streamlit as st
import asyncio, tempfile, urllib.parse, requests, pandas as pd
import plotly.express as px
from datetime import datetime, timedelta
from groq import Groq
from tavily import TavilyClient
import google.generativeai as genai
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
import edge_tts

PROXY_URL = "https://groq-proxy.irantradedesk.workers.dev/"
groq_client = Groq(api_key=st.secrets["GROQ_API_KEY"], base_url=PROXY_URL)
tavily_client = TavilyClient(api_key=st.secrets["TAVILY_API_KEY"])
genai.configure(api_key=st.secrets["GEMINI_API_KEY"])

def ask_groq(messages, model_name="openai/gpt-oss-120b"):
    r = groq_client.chat.completions.create(messages=messages, model=model_name)
    return r.choices[0].message.content

def ask_gemini(messages, model_name="gemini-3.8-flash"):
    system_prompt, history = "", []
    for msg in messages:
        if msg["role"] == "system": system_prompt = msg["content"]; continue
        role = "user" if msg["role"] == "user" else "model"
        history.append({"role": role, "parts": [msg["content"]]})
    model = genai.GenerativeModel(model_name, system_instruction=system_prompt) if system_prompt else genai.GenerativeModel(model_name)
    return model.generate_content(history).text

def ask(messages, provider):
    return ask_groq(messages) if provider == "Groq (سریع)" else ask_gemini(messages)

def search_web(q):
    try:
        r = tavily_client.search(query=q, max_results=5)
        return "\n".join(f"- {x['title']}: {x['content']}" for x in r.get("results", []))
    except: return ""

def get_forex():
    try:
        r = requests.get("https://api.frankfurter.app/latest?from=USD&to=EUR,GBP,AED,CNY,JPY,TRY", timeout=10)
        return r.json().get("rates", {})
    except Exception as e: return {"error": str(e)}

def get_crypto():
    try:
        url = "https://api.coingecko.com/api/v3/simple/price?ids=bitcoin,ethereum,tether&vs_currencies=usd"
        return requests.get(url, timeout=10).json()
    except Exception as e: return {"error": str(e)}

W_CODES = {0:"صاف ☀️",1:"عمدتاً صاف 🌤️",2:"نیمه ابری ⛅",3:"ابری ☁️",45:"مه 🌫️",51:"نم‌نم 🌦️",61:"باران 🌧️",71:"برف 🌨️",80:"رگبار 🌦️",95:"رعد ⛈️"}

def get_weather(city):
    try:
        geo = requests.get(f"https://geocoding-api.open-meteo.com/v1/search?name={urllib.parse.quote(city)}&count=1&language=fa", timeout=10).json()
        if not geo.get("results"): return {"error": "شهر پیدا نشد"}
        loc = geo["results"][0]
        w = requests.get(f"https://api.open-meteo.com/v1/forecast?latitude={loc['latitude']}&longitude={loc['longitude']}&current=temperature_2m,relative_humidity_2m,wind_speed_10m,weather_code", timeout=10).json()
        return {"city": loc.get("name", city), "country": loc.get("country", ""), "data": w.get("current", {})}
    except Exception as e: return {"error": str(e)}

def read_file(uploaded):
    name = uploaded.name.lower()
    try:
        if name.endswith(".pdf"): return "\n".join((p.extract_text() or "") for p in PdfReader(uploaded).pages)
        if name.endswith(".docx"): return "\n".join(p.text for p in Document(uploaded).paragraphs)
        if name.endswith((".xlsx", ".xls")): return pd.read_excel(uploaded).to_string()
        if name.endswith(".csv"): return pd.read_csv(uploaded).to_string()
        if name.endswith((".txt", ".md")): return uploaded.read().decode("utf-8", errors="ignore")
    except: pass
    return ""

def build_pptx(title, slides_text):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = "Iran Trade Services"
    # Content slides
    blocks = [b.strip() for b in slides_text.split("---") if b.strip()]
    for block in blocks:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines: continue
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = lines[0]
        tf = slide.placeholders[1].text_frame
        tf.text = lines[1] if len(lines) > 1 else ""
        for line in lines[2:]:
            tf.add_paragraph().text = line
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
    prs.save(tmp.name)
    return tmp.name

async def _tts(text, voice, path):
    await edge_tts.Communicate(text, voice).save(path)

def text_to_speech(text, lang="فارسی"):
    voices = {"فارسی": "fa-IR-DilaraNeural", "انگلیسی": "en-US-AriaNeural", "عربی": "ar-SA-ZariyahNeural"}
    with tempfile.NamedTemporaryFile(delete=False, suffix=".mp3") as f: path = f.name
    asyncio.run(_tts(text, voices.get(lang, voices["فارسی"]), path))
    return path

def generate_image(prompt, w=1024, h=1024):
    return f"https://image.pollinations.ai/prompt/{urllib.parse.quote(prompt)}?width={w}&height={h}&nologo=true"

# =====================================================
# UI
# =====================================================
st.set_page_config(page_title="Iran Trade Services", page_icon="🌍", layout="wide")

# --- احراز هویت ---
if "authenticated" not in st.session_state:
    st.session_state.authenticated = False

if not st.session_state.authenticated:
    st.title("🔐 Iran Trade Services")
    st.caption("لطفاً رمز عبور را وارد کنید")
    pwd = st.text_input("رمز عبور:", type="password")
    if st.button("ورود", type="primary"):
        if pwd == st.secrets.get("APP_PASSWORD", ""):
            st.session_state.authenticated = True
            st.rerun()
        else:
            st.error("❌ رمز اشتباه است")
    st.stop()

# --- UI اصلی ---

st.title("🌍 Iran Trade Services")
st.caption("دستیار هوشمند تجارت بین‌الملل")

with st.sidebar:
    st.header("⚙️ تنظیمات")
    provider = st.radio("موتور هوش مصنوعی:", ("Groq (سریع)", "Gemini (پایدار)"))
    st.markdown("---")
    mode = st.radio("حالت کار:", (
        "💬 چت آزاد", "📱 تولید محتوا", "🌐 ترجمه حرفه‌ای", "📄 تحلیل فایل",
        "📊 ساخت پاورپوینت", "📈 داده‌های مالی", "🌤️ آب‌وهوا", "📰 اخبار",
        "🖼️ تولید تصویر", "🔊 متن به صدا", "🎤 صدا به متن",
        "📅 تقویم و یادآوری", "📊 گانت چارت", "✉️ پیش‌نویس ایمیل"
    ))

# ۱. چت آزاد
if mode == "💬 چت آزاد":
    if "messages" not in st.session_state:
        st.session_state.messages = [{"role": "system", "content": "تو دستیار شخصی من هستی. فارسی، دقیق و مفید پاسخ بده."}]
    if "display" not in st.session_state: st.session_state.display = []
    for msg in st.session_state.display:
        with st.chat_message(msg["role"]): st.write(msg["content"])
    user_input = st.chat_input("سوالت رو بنویس...")
    if user_input:
        st.session_state.display.append({"role": "user", "content": user_input})
        with st.chat_message("user"): st.write(user_input)
        with st.chat_message("assistant"):
            with st.spinner("در حال فکر کردن..."):
                search = search_web(user_input)
                full_msg = f"سوال: {user_input}\n\nنتایج جستجو:\n{search}\n\nپاسخ فارسی بده." if search else user_input
                st.session_state.messages.append({"role": "user", "content": full_msg})
                try: answer = ask(st.session_state.messages, provider)
                except Exception as e: answer = f"خطا: {e}"
                st.session_state.messages.append({"role": "assistant", "content": answer})
                st.session_state.display.append({"role": "assistant", "content": answer})
                st.write(answer)

# ۲. تولید محتوا
elif mode == "📱 تولید محتوا":
    st.subheader("📱 استودیوی تولید محتوا")
    c1, c2 = st.columns(2)
    with c1:
        platform = st.selectbox("پلتفرم:", ["LinkedIn", "Instagram", "Twitter/X", "Telegram", "Facebook"])
        content_type = st.selectbox("نوع محتوا:", ["پست کامل", "کپشن کوتاه", "تقویم هفتگی", "۱۰ ایده", "سناریوی Reels", "هشتگ‌ها"])
    with c2:
        tone = st.selectbox("لحن:", ["حرفه‌ای", "دوستانه", "آموزشی", "تبلیغاتی", "الهام‌بخش"])
        language = st.selectbox("زبان:", ["فارسی", "انگلیسی", "دوزبانه"])
    topic = st.text_area("موضوع:", height=100)
    if st.button("✨ تولید محتوا", type="primary"):
        if not topic.strip(): st.warning("موضوع رو وارد کن.")
        else:
            with st.spinner("در حال تولید..."):
                p = f"محتوای {content_type} برای {platform}، لحن {tone}، زبان {language}، موضوع: {topic}"
                try:
                    r = ask([{"role": "system", "content": "تو متخصص محتوای شبکه‌های اجتماعی هستی."}, {"role": "user", "content": p}], provider)
                    st.markdown("### 📄 محتوا:"); st.markdown(r)
                except Exception as e: st.error(f"خطا: {e}")

# ۳. ترجمه
elif mode == "🌐 ترجمه حرفه‌ای":
    st.subheader("🌐 ترجمه حرفه‌ای")
    c1, c2 = st.columns(2)
    with c1:
        src = st.selectbox("مبدأ:", ["تشخیص خودکار", "فارسی", "انگلیسی", "عربی", "چینی", "ترکی", "روسی"])
        ctx = st.selectbox("زمینه:", ["عمومی", "تجاری", "حقوقی و قرارداد", "فنی", "بازاریابی", "مکاتبات رسمی"])
    with c2:
        tgt = st.selectbox("مقصد:", ["انگلیسی", "فارسی", "عربی", "چینی", "ترکی", "روسی"])
        style = st.selectbox("سبک:", ["دقیق", "روان", "رسمی", "محاوره‌ای"])
    text_in = st.text_area("متن:", height=200)
    if st.button("🌐 ترجمه", type="primary"):
        if not text_in.strip(): st.warning("متن رو وارد کن.")
        else:
            with st.spinner("در حال ترجمه..."):
                p = f"از {src} به {tgt}، زمینه {ctx}، سبک {style}:\n\n{text_in}"
                try:
                    r = ask([{"role": "system", "content": "مترجم حرفه‌ای هستی."}, {"role": "user", "content": p}], provider)
                    st.markdown(r)
                except Exception as e: st.error(f"خطا: {e}")

# ۴. تحلیل فایل
elif mode == "📄 تحلیل فایل":
    st.subheader("📄 تحلیل هوشمند فایل")
    uploaded = st.file_uploader("فایل:", type=["pdf", "docx", "xlsx", "csv", "txt"])
    if uploaded:
        content = read_file(uploaded)
        if content:
            st.success(f"✅ {len(content):,} کاراکتر خوانده شد")
            with st.expander("پیش‌نمایش"): st.text(content[:1000] + ("..." if len(content) > 1000 else ""))
            action = st.selectbox("عملیات:", ["خلاصه", "نکات کلیدی", "تحلیل قرارداد", "ترجمه به انگلیسی", "استخراج اطلاعات"])
            if st.button("🚀 اجرا", type="primary"):
                with st.spinner("در حال تحلیل..."):
                    try:
                        r = ask([{"role": "system", "content": "تحلیلگر اسناد تجاری هستی."}, {"role": "user", "content": f"{action}:\n\n{content[:15000]}"}], provider)
                        st.markdown(r)
                    except Exception as e: st.error(f"خطا: {e}")

# ۵. پاورپوینت
elif mode == "📊 ساخت پاورپوینت":
    st.subheader("📊 ساخت پاورپوینت")
    title = st.text_input("عنوان ارائه:", value="گزارش Iran Trade Services")
    topic = st.text_area("موضوع و محتوای ارائه:", height=200, placeholder="مثلاً: گزارش عملکرد صادرات زعفران در سال ۱۴۰۴، شامل بازارها، رقبا، استراتژی...")
    n_slides = st.slider("تعداد اسلاید:", 3, 15, 6)
    if st.button("📊 ساخت پاورپوینت", type="primary"):
        if not topic.strip(): st.warning("موضوع رو وارد کن.")
        else:
            with st.spinner("در حال ساخت..."):
                try:
                    p = f"یک ارائه {n_slides} اسلایدی درباره موضوع زیر بساز. اسلایدها رو با '---' جدا کن. هر اسلاید: خط اول = عنوان، خطوط بعدی = محتوا.\n\nموضوع: {topic}"
                    slides_text = ask([{"role": "system", "content": "تو متخصص ساخت ارائه هستی."}, {"role": "user", "content": p}], provider)
                    path = build_pptx(title, slides_text)
                    with open(path, "rb") as f: pptx_data = f.read()
                    st.success("✅ پاورپوینت آماده شد!")
                    st.download_button("📥 دانلود PPTX", pptx_data, file_name="presentation.pptx", mime="application/vnd.openxmlformats-officedocument.presentationml.presentation")
                    with st.expander("پیش‌نمایش محتوا"): st.markdown(slides_text)
                except Exception as e: st.error(f"خطا: {e}")

# ۶. داده‌های مالی
elif mode == "📈 داده‌های مالی":
    st.subheader("📈 داده‌های زنده")
    tab1, tab2 = st.tabs(["💱 نرخ ارز", "💰 کریپتو"])
    with tab1:
        with st.spinner("در حال دریافت..."):
            fx = get_forex()
            if "error" in fx: st.error(fx["error"])
            else: st.dataframe(pd.DataFrame([{"ارز": k, "نرخ": v} for k, v in fx.items()]), use_container_width=True)
    with tab2:
        with st.spinner("در حال دریافت..."):
            c = get_crypto()
            if "error" in c: st.error(c["error"])
            else: st.dataframe(pd.DataFrame([{"ارز": k, "دلار": f"${v['usd']:,.2f}"} for k, v in c.items()]), use_container_width=True)

# ۷. آب‌وهوا
elif mode == "🌤️ آب‌وهوا":
    st.subheader("🌤️ آب‌وهوای مقاصد تجاری")
    city = st.text_input("نام شهر (انگلیسی):", placeholder="Dubai, Istanbul, Shanghai, Hamburg")
    if st.button("🔍 جستجو"):
        if city:
            with st.spinner("در حال دریافت..."):
                w = get_weather(city)
                if "error" in w: st.error(w["error"])
                else:
                    d = w["data"]
                    st.markdown(f"### 📍 {w['city']}, {w['country']}")
                    c1, c2, c3 = st.columns(3)
                    c1.metric("🌡️ دما", f"{d.get('temperature_2m', '-')}°C")
                    c2.metric("💧 رطوبت", f"{d.get('relative_humidity_2m', '-')}%")
                    c3.metric("💨 باد", f"{d.get('wind_speed_10m', '-')} km/h")
                    st.info(W_CODES.get(d.get('weather_code'), "نامشخص"))

# ۸. اخبار
elif mode == "📰 اخبار":
    st.subheader("📰 اخبار تجاری")
    query = st.text_input("موضوع:", value="صادرات و واردات ایران")
    count = st.slider("تعداد:", 3, 10, 5)
    if st.button("📰 دریافت", type="primary"):
        with st.spinner("در حال جستجو..."):
            try:
                res = tavily_client.search(query=query, max_results=count, topic="news")
                for r in res.get("results", []):
                    st.markdown(f"### 📌 {r.get('title', '')}")
                    st.write(r.get("content", ""))
                    st.caption(f"[منبع]({r.get('url', '#')})"); st.markdown("---")
            except Exception as e: st.error(f"خطا: {e}")

# ۹. تولید تصویر
elif mode == "🖼️ تولید تصویر":
    st.subheader("🖼️ تولید تصویر تبلیغاتی")
    prompt = st.text_area("توصیف (انگلیسی):", placeholder="A modern logistics warehouse, professional photography, 4k", height=100)
    c1, c2 = st.columns(2)
    with c1: w = st.selectbox("عرض:", [512, 768, 1024, 1280], index=2)
    with c2: h = st.selectbox("ارتفاع:", [512, 768, 1024, 1280], index=2)
    if st.button("🎨 تولید تصویر", type="primary"):
        if not prompt.strip(): st.warning("توصیف رو وارد کن.")
        else:
            url = generate_image(prompt, w, h)
            st.success("✅ آماده شد!")
            st.image(url, caption=prompt, use_container_width=True)
            st.caption(f"[دانلود]({url})")

# ۱۰. متن به صدا
elif mode == "🔊 متن به صدا":
    st.subheader("🔊 تبدیل متن به صدا")
    text = st.text_area("متن:", height=200)
    lang = st.selectbox("زبان:", ["فارسی", "انگلیسی", "عربی"])
    if st.button("🎤 تولید صدا", type="primary"):
        if not text.strip(): st.warning("متن رو وارد کن.")
        else:
            with st.spinner("در حال تولید..."):
                try:
                    path = text_to_speech(text, lang)
                    with open(path, "rb") as f: audio = f.read()
                    st.audio(audio, format="audio/mp3")
                    st.download_button("📥 دانلود", audio, file_name="voice.mp3", mime="audio/mp3")
                except Exception as e: st.error(f"خطا: {e}")

# ۱۱. صدا به متن
elif mode == "🎤 صدا به متن":
    st.subheader("🎤 تبدیل صدا به متن")
    audio_file = st.file_uploader("فایل صوتی:", type=["mp3", "wav", "m4a", "ogg", "webm"])
    if audio_file:
        if st.button("📝 تبدیل به متن", type="primary"):
            with st.spinner("در حال تبدیل..."):
                try:
                    tr = groq_client.audio.transcriptions.create(
                        file=(audio_file.name, audio_file.read()),
                        model="whisper-large-v3-turbo",
                        response_format="text"
                    )
                    st.success("✅ متن استخراج شد!")
                    st.text_area("متن:", tr, height=250)
                except Exception as e: st.error(f"خطا: {e}")

# ۱۲. تقویم
elif mode == "📅 تقویم و یادآوری":
    st.subheader("📅 تقویم و یادآوری")
    if "events" not in st.session_state: st.session_state.events = []
    with st.form("add_event"):
        st.write("**افزودن رویداد جدید**")
        c1, c2 = st.columns(2)
        with c1:
            ev_title = st.text_input("عنوان:")
            ev_date = st.date_input("تاریخ:")
        with c2:
            ev_time = st.time_input("ساعت:")
            ev_note = st.text_input("یادداشت:")
        if st.form_submit_button("➕ افزودن"):
            if ev_title.strip():
                st.session_state.events.append({"عنوان": ev_title, "تاریخ": str(ev_date), "ساعت": str(ev_time), "یادداشت": ev_note})
                st.success(f"✅ رویداد «{ev_title}» اضافه شد.")
    if st.session_state.events:
        st.markdown("### 📋 رویدادهای ثبت‌شده")
        st.dataframe(pd.DataFrame(st.session_state.events), use_container_width=True)
        if st.button("🗑️ پاک کردن همه"):
            st.session_state.events = []; st.rerun()

# ۱۳. گانت چارت
elif mode == "📊 گانت چارت":
    st.subheader("📊 گانت چارت برنامه‌ریزی پروژه")
    st.write("پروژه رو توصیف کن، من گانت چارت می‌سازم.")
    project_desc = st.text_area("توصیف پروژه:", height=150, placeholder="مثلاً: پروژه صادرات ۱۰ تن زعفران به چین. مراحل: تحقیق بازار، مذاکره با خریدار، عقد قرارداد، بسته‌بندی، حمل، تحویل. مدت کل: ۶ ماه")
    if st.button("📊 ساخت گانت چارت", type="primary"):
        if not project_desc.strip(): st.warning("توصیف رو وارد کن.")
        else:
            with st.spinner("در حال ساخت..."):
                p = f"""پروژه زیر رو تحلیل کن و به صورت جدول CSV با ستون‌های: "Task","Start","Finish","Resource" استخراج کن.
تاریخ‌ها به فرمت YYYY-MM-DD باشن. امروز: {datetime.now().strftime('%Y-%m-%d')}.

پروژه: {project_desc}

فقط CSV بده، بدون توضیح اضافه."""
                try:
                    csv_text = ask([{"role": "system", "content": "تو مدیر پروژه حرفه‌ای هستی."}, {"role": "user", "content": p}], provider)
                    import io
                    csv_clean = csv_text.replace("```csv", "").replace("```", "").strip()
                    df = pd.read_csv(io.StringIO(csv_clean))
                    fig = px.timeline(df, x_start="Start", x_end="Finish", y="Task", color="Resource", title="گانت چارت پروژه")
                    fig.update_yaxes(autorange="reversed")
                    st.plotly_chart(fig, use_container_width=True)
                    st.dataframe(df, use_container_width=True)
                except Exception as e: st.error(f"خطا: {e}")
                st.caption("💡 این گانت چارت تعاملیه — می‌تونی زوم کنی و جزئیات رو ببینی.")

# ۱۴. پیش‌نویس ایمیل
elif mode == "✉️ پیش‌نویس ایمیل":
    st.subheader("✉️ پیش‌نویس ایمیل تجاری")
    st.info("💡 این بخش فقط پیش‌نویس ایمیل رو آماده می‌کنه. برای ارسال مستقیم، باید Gmail API رو وصل کنیم.")
    c1, c2 = st.columns(2)
    with c1:
        recipient = st.text_input("گیرنده:", placeholder="manager@company.com")
        subject = st.text_input("موضوع ایمیل:", placeholder="Proposal for Saffron Export")
    with c2:
        email_lang = st.selectbox("زبان:", ["انگلیسی", "فارسی", "عربی"])
        email_tone = st.selectbox("لحن:", ["رسمی", "دوستانه", "تجاری"])
    purpose = st.text_area("هدف ایمیل:", height=150, placeholder="مثلاً: معرفی خدمات Iran Trade Services و پیشنهاد همکاری در واردات زعفران")
    if st.button("✉️ نوشتن پیش‌نویس", type="primary"):
        if not purpose.strip(): st.warning("هدف رو وارد کن.")
        else:
            with st.spinner("در حال نوشتن..."):
                p = f"""یک ایمیل {email_tone} به زبان {email_lang} بنویس.
گیرنده: {recipient}
موضوع: {subject}
هدف: {purpose}

ایمیل باید حرفه‌ای، واضح و با CTA مناسب باشه."""
                try:
                    r = ask([{"role": "system", "content": "تو متخصص نگارش ایمیل تجاری هستی."}, {"role": "user", "content": p}], provider)
                    st.markdown("### 📧 پیش‌نویس:")
                    st.markdown(r)
                    st.caption("💡 این متن رو کپی کن و توی Gmail بفرست.")
                except Exception as e: st.error(f"خطا: {e}")
